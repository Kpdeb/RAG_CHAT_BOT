import os
import shutil
import requests

from fastapi import (
    FastAPI,
    UploadFile,
    File,
    HTTPException
)

from fastapi.middleware.cors import CORSMiddleware

from pydantic import BaseModel

from website_loader import load_website

from rag_pipeline import ask_question

from chroma import (
    add_documents,
    get_document_count,
    get_chunk_count,
    get_database_stats,
    get_documents,
    delete_document
)

from chat_db import (
    init_db,
    create_chat,
    get_chats,
    add_message,
    get_messages,
    update_chat_title,
    delete_chat
)

from langchain_community.document_loaders import (
    PyPDFLoader
)

from langchain_core.documents import Document


# =========================================================
# BASE DIRECTORY
# =========================================================

BASE_DIR = os.path.dirname(
    os.path.abspath(__file__)
)


# =========================================================
# APP
# =========================================================

app = FastAPI(
    title="AI Knowledge Assistant API",
    version="1.0.0"
)


# =========================================================
# INITIALIZE CHAT DATABASE
# =========================================================

init_db()


# =========================================================
# CORS
# =========================================================

FRONTEND_URL = os.getenv(
    "FRONTEND_URL",
    "http://localhost:3000"
)

app.add_middleware(
    CORSMiddleware,
    allow_origins=[FRONTEND_URL],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)


# =========================================================
# UPLOAD DIRECTORY
# =========================================================

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
UPLOAD_DIR = os.path.join(BASE_DIR, "document_loader")

os.makedirs(
    UPLOAD_DIR,
    exist_ok=True
)


# =========================================================
# RECENT ACTIVITY
# =========================================================

recent_activities = []


def add_activity(
    activity_type: str,
    message: str,
    details: str = ""
):

    from datetime import datetime

    recent_activities.insert(
        0,
        {
            "type": activity_type,
            "message": message,
            "details": details,
            "timestamp": datetime.now().isoformat()
        }
    )

    # Keep only latest 10 activities
    del recent_activities[10:]


# =========================================================
# CHAT REQUEST
# =========================================================

class ChatRequest(BaseModel):

    question: str
    chat_id: int


# =========================================================
# WEBSITE REQUEST
# =========================================================

class WebsiteRequest(BaseModel):

    url: str


# =========================================================
# ROOT
# =========================================================

@app.get("/")
def root():

    return {
        "message": "AI Knowledge Assistant API is running"
    }


# =========================================================
# HEALTH
# =========================================================

@app.get("/health")
def health():

    return {
        "status": "ok"
    }


# =========================================================
# DOCUMENT COUNT
# =========================================================

@app.get("/documents/count")
def document_count():

    return {
        "count": get_document_count()
    }


# =========================================================
# DOCUMENT STATS
# =========================================================

@app.get("/documents/stats")
def document_stats():

    return get_database_stats()


# =========================================================
# RECENT ACTIVITY
# =========================================================

@app.get("/activity")
def get_activity():

    return {
        "activities": recent_activities
    }


# =========================================================
# GET DOCUMENTS
# =========================================================

@app.get("/documents")
def documents():

    return {
        "documents": get_documents()
    }


# =========================================================
# DELETE DOCUMENT
# =========================================================

@app.delete("/documents/{filename}")
def delete_document_endpoint(
    filename: str
):

    try:

        deleted_chunks = delete_document(
            filename
        )

        if deleted_chunks == 0:

            raise HTTPException(
                status_code=404,
                detail="Document not found."
            )

        # -------------------------------------------------
        # Delete original uploaded file
        # -------------------------------------------------

        file_path = os.path.join(
            UPLOAD_DIR,
            filename
        )

        if os.path.exists(file_path):

            os.remove(file_path)

        # -------------------------------------------------
        # Record activity
        # -------------------------------------------------

        add_activity(
            "delete",
            f"{filename} deleted",
            f"{deleted_chunks} chunks removed from the knowledge base."
        )

        # -------------------------------------------------
        # Updated stats
        # -------------------------------------------------

        stats = get_database_stats()

        return {
            "message": "Document deleted successfully.",
            "filename": filename,
            "chunks_deleted": deleted_chunks,
            "total_documents": stats["documents"],
            "total_chunks": stats["chunks"]
        }

    except HTTPException:

        raise

    except Exception as e:

        print(
            "DELETE ERROR:",
            repr(e)
        )

        raise HTTPException(
            status_code=500,
            detail=str(e)
        )


# =========================================================
# WEBSITE INGESTION
# =========================================================

@app.post("/websites")
def add_website(
    request: WebsiteRequest
):

    url = request.url.strip()

    # -----------------------------------------------------
    # Validate URL
    # -----------------------------------------------------

    if not url:

        raise HTTPException(
            status_code=400,
            detail="Website URL is required."
        )

    # -----------------------------------------------------
    # Add HTTPS automatically
    # -----------------------------------------------------

    if not url.startswith(
        ("http://", "https://")
    ):

        url = "https://" + url

    try:

        print(
            f"Website ingestion started: {url}"
        )

        # =================================================
        # LOAD WEBSITE
        # =================================================

        documents = load_website(
            url
        )

        if not documents:

            raise HTTPException(
                status_code=400,
                detail=(
                    "Could not extract readable "
                    "content from this website."
                )
            )

        print(
            f"Loaded {len(documents)} website document(s)"
        )

        # =================================================
        # REMOVE OLD VERSION
        # =================================================

        try:

            deleted_chunks = delete_document(
                url
            )

            if deleted_chunks > 0:

                print(
                    f"Removed {deleted_chunks} old website chunks"
                )

        except Exception as delete_error:

            print(
                "OLD WEBSITE DELETE WARNING:",
                repr(delete_error)
            )

        # =================================================
        # ADD WEBSITE TO CHROMA
        # =================================================

        chunks_added = add_documents(
            documents
        )

        print(
            f"Added {chunks_added} website chunks to Chroma"
        )

        # =================================================
        # RECORD ACTIVITY
        # =================================================

        add_activity(
            "website",
            "Website added",
            f"{url} added with {chunks_added} chunks."
        )

        # =================================================
        # UPDATED STATS
        # =================================================

        stats = get_database_stats()

        return {
            "message": "Website added successfully.",
            "url": url,
            "documents_loaded": len(documents),
            "chunks_added": chunks_added,
            "total_documents": stats["documents"],
            "total_chunks": stats["chunks"]
        }

    except requests.exceptions.RequestException as e:

        print(
            "WEBSITE REQUEST ERROR:",
            repr(e)
        )

        raise HTTPException(
            status_code=400,
            detail=(
                "Could not access the website. "
                "Please check the URL and make sure "
                "the website is publicly accessible."
            )
        )

    except ValueError as e:

        print(
            "WEBSITE VALUE ERROR:",
            repr(e)
        )

        raise HTTPException(
            status_code=400,
            detail=str(e)
        )

    except HTTPException:

        raise

    except Exception as e:

        print(
            "WEBSITE INGESTION ERROR:",
            repr(e)
        )

        raise HTTPException(
            status_code=500,
            detail=(
                "Failed to process the website."
            )
        )


# =========================================================
# CHAT HISTORY
# =========================================================


# ---------------------------------------------------------
# CREATE NEW CHAT
# ---------------------------------------------------------

@app.post("/chats")
def create_new_chat():

    try:

        chat_id = create_chat(
            title="New Chat"
        )

        return {
            "chat_id": chat_id,
            "title": "New Chat"
        }

    except Exception as e:

        print(
            "CREATE CHAT ERROR:",
            repr(e)
        )

        raise HTTPException(
            status_code=500,
            detail=str(e)
        )


# ---------------------------------------------------------
# GET ALL CHATS
# ---------------------------------------------------------

@app.get("/chats")
def get_all_chats():

    try:

        chats = get_chats()

        return {
            "chats": chats
        }

    except Exception as e:

        print(
            "GET CHATS ERROR:",
            repr(e)
        )

        raise HTTPException(
            status_code=500,
            detail=str(e)
        )


# ---------------------------------------------------------
# GET CHAT MESSAGES
# ---------------------------------------------------------

@app.get("/chats/{chat_id}/messages")
def chat_messages(
    chat_id: int
):

    try:

        messages = get_messages(
            chat_id
        )

        return {
            "chat_id": chat_id,
            "messages": messages
        }

    except Exception as e:

        print(
            "GET MESSAGES ERROR:",
            repr(e)
        )

        raise HTTPException(
            status_code=500,
            detail=str(e)
        )


# ---------------------------------------------------------
# GET CHAT
# ---------------------------------------------------------

@app.get("/chats/{chat_id}")
def get_chat_messages(
    chat_id: int
):

    try:

        messages = get_messages(
            chat_id
        )

        return {
            "chat_id": chat_id,
            "messages": messages
        }

    except Exception as e:

        print(
            "GET CHAT ERROR:",
            repr(e)
        )

        raise HTTPException(
            status_code=500,
            detail=str(e)
        )


# ---------------------------------------------------------
# DELETE CHAT
# ---------------------------------------------------------

@app.delete("/chats/{chat_id}")
def remove_chat(
    chat_id: int
):

    try:

        delete_chat(
            chat_id
        )

        return {
            "message": "Chat deleted successfully."
        }

    except Exception as e:

        print(
            "DELETE CHAT ERROR:",
            repr(e)
        )

        raise HTTPException(
            status_code=500,
            detail=str(e)
        )


# =========================================================
# CHAT
# =========================================================

@app.post("/chat")
def chat(
    request: ChatRequest
):

    question = request.question.strip()

    # -----------------------------------------------------
    # Empty question
    # -----------------------------------------------------

    if not question:

        return {
            "answer": "Please enter a question.",
            "sources": []
        }

    try:

        # =================================================
        # CHECK CHAT EXISTS
        # =================================================

        chats = get_chats()

        current_chat = next(
            (
                chat
                for chat in chats
                if chat["id"] == request.chat_id
            ),
            None
        )

        if current_chat is None:

            raise HTTPException(
                status_code=404,
                detail="Chat session not found."
            )

        # =================================================
        # SAVE USER MESSAGE
        # =================================================

        add_message(
            chat_id=request.chat_id,
            role="user",
            content=question,
            sources=[]
        )

        # =================================================
        # GENERATE CHAT TITLE
        # =================================================

        if current_chat["title"] == "New Chat":

            title = question.strip()

            if len(title) > 40:

                title = (
                    title[:40].rstrip()
                    + "..."
                )

            update_chat_title(
                chat_id=request.chat_id,
                title=title
            )

        # =================================================
        # ASK RAG PIPELINE
        # =================================================

        result = ask_question(
            question
        )

        # Make sure result is a dictionary
        if not isinstance(result, dict):

            result = {
                "answer": str(result),
                "sources": []
            }

        answer = result.get(
            "answer",
            "I could not find an answer."
        )

        sources = result.get(
            "sources",
            []
        )

        # =================================================
        # SAVE ASSISTANT RESPONSE
        # =================================================

        add_message(
            chat_id=request.chat_id,
            role="assistant",
            content=answer,
            sources=sources
        )

        # =================================================
        # RETURN RESPONSE
        # =================================================

        return {
            "answer": answer,
            "sources": sources
        }

    except HTTPException:

        raise

    except Exception as e:

        print(
            "CHAT ERROR:",
            repr(e)
        )

        return {
            "answer": "",
            "sources": [],
            "error": str(e)
        }


# =========================================================
# DOCUMENT PARSER
# =========================================================

def load_document(
    file_path: str,
    filename: str
):

    extension = os.path.splitext(
        filename
    )[1].lower()

    # =====================================================
    # PDF
    # =====================================================

    if extension == ".pdf":

        loader = PyPDFLoader(
            file_path
        )

        documents = loader.load()

        for doc in documents:

            doc.metadata["file_name"] = filename
            doc.metadata["source_type"] = "pdf"

        return documents

    # =====================================================
    # TXT
    # =====================================================

    elif extension == ".txt":

        with open(
            file_path,
            "r",
            encoding="utf-8"
        ) as f:

            text = f.read()

        return [
            Document(
                page_content=text,
                metadata={
                    "file_name": filename,
                    "source_type": "txt"
                }
            )
        ]

    # =====================================================
    # MARKDOWN
    # =====================================================

    elif extension == ".md":

        with open(
            file_path,
            "r",
            encoding="utf-8"
        ) as f:

            text = f.read()

        return [
            Document(
                page_content=text,
                metadata={
                    "file_name": filename,
                    "source_type": "markdown"
                }
            )
        ]

    # =====================================================
    # DOCX
    # =====================================================

    elif extension == ".docx":

        from docx import Document as DocxDocument

        docx_file = DocxDocument(
            file_path
        )

        paragraphs = []

        for paragraph in docx_file.paragraphs:

            text = paragraph.text.strip()

            if text:

                paragraphs.append(
                    text
                )

        full_text = "\n".join(
            paragraphs
        )

        return [
            Document(
                page_content=full_text,
                metadata={
                    "file_name": filename,
                    "source_type": "docx"
                }
            )
        ]

    # =====================================================
    # CSV
    # =====================================================

    elif extension == ".csv":

        import pandas as pd

        df = pd.read_csv(
            file_path
        )

        text = df.to_string(
            index=False
        )

        return [
            Document(
                page_content=text,
                metadata={
                    "file_name": filename,
                    "source_type": "csv"
                }
            )
        ]

    # =====================================================
    # UNSUPPORTED
    # =====================================================

    else:

        raise ValueError(
            f"Unsupported file type: {extension}"
        )


# =========================================================
# UPLOAD DOCUMENT
# =========================================================

@app.post("/upload")
async def upload_document(
    file: UploadFile = File(...)
):

    filename = file.filename

    if not filename:

        raise HTTPException(
            status_code=400,
            detail="No filename provided."
        )

    # =====================================================
    # ALLOWED FILES
    # =====================================================

    allowed_extensions = {
        ".pdf",
        ".docx",
        ".txt",
        ".md",
        ".csv"
    }

    extension = os.path.splitext(
        filename
    )[1].lower()

    if extension not in allowed_extensions:

        raise HTTPException(
            status_code=400,
            detail=(
                "Unsupported file type. "
                "Allowed: PDF, DOCX, TXT, MD, CSV"
            )
        )

    # =====================================================
    # SAVE FILE
    # =====================================================

    file_path = os.path.join(
        UPLOAD_DIR,
        filename
    )

    try:

        with open(
            file_path,
            "wb"
        ) as buffer:

            shutil.copyfileobj(
                file.file,
                buffer
            )

        print(
            f"Uploaded: {filename}"
        )

        # =================================================
        # LOAD DOCUMENT
        # =================================================

        documents = load_document(
            file_path,
            filename
        )

        if not documents:

            raise ValueError(
                "No readable content found."
            )

        print(
            f"Loaded {len(documents)} document sections"
        )

        # =================================================
        # ADD TO CHROMA
        # =================================================

        chunks_added = add_documents(
            documents
        )

        print(
            f"Added {chunks_added} chunks to Chroma"
        )

        # =================================================
        # RECORD ACTIVITY
        # =================================================

        add_activity(
            "upload",
            f"{filename} uploaded",
            f"{chunks_added} chunks added to the knowledge base."
        )

        # =================================================
        # UPDATED STATS
        # =================================================

        stats = get_database_stats()

        return {
            "message": "Document uploaded successfully.",
            "filename": filename,
            "type": extension,
            "documents_loaded": len(documents),
            "chunks_added": chunks_added,
            "total_documents": stats["documents"],
            "total_chunks": stats["chunks"]
        }

    except Exception as e:

        print(
            "UPLOAD ERROR:",
            repr(e)
        )

        # Remove partially saved file
        if os.path.exists(file_path):

            os.remove(file_path)

        raise HTTPException(
            status_code=500,
            detail=str(e)
        )

    finally:

        await file.close()